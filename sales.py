import pandas as pd
import sqlite3
from openpyxl import load_workbook
from tkinter.filedialog import askopenfilename
from tkinter import messagebox

def read_sales():
    print("read_sales() called")


    # Select Excel File
    file_path = askopenfilename(title="Seleccionar archivo Excel de Ventas", filetypes=[("Archivos Excel", "*.xlsx *.xls")])

    # Load Excel
    excel_file = pd.ExcelFile(file_path)
    results = []
    sheet_names = excel_file.sheet_names
    sheet_name = r"Informe de ventas e IVA"
    wb = load_workbook(file_path, data_only=True)
    if sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]  # Activate the specific sheet
    else:
        print(f"Sheet '{sheet_name}' not found in the workbook.")
        sheet = None

    # Extract visible rows (non-hidden rows)
    visible_rows = []
    for row in sheet.iter_rows():
        if not sheet.row_dimensions[row[0].row].hidden:  # Check if the row is hidden
            visible_rows.append([cell.value for cell in row])

    # Convert the visible rows into a DataFrame
    data = pd.DataFrame(visible_rows)
    stop_word = r"Total"
    cols = list(data.columns)
    first_row = data.iloc[4][0]
    month_year = ' '.join(first_row.split(' ')[-2:])
    month = month_year.split('/')[1]
    year = month_year.split('/')[2].split(')')[0]
    data['Month'] = month
    data['Year'] = year
    index: int
    flag_write = 0
    for index, row in data.iterrows():
        if str(row[0]).__contains__("País de origen"):
            flag_write = 1
        if str(row[0]).__contains__("Total") or str(row[0]).strip() == "":
            flag_write = 0
        if flag_write == 1 and not str(row[0]).__contains__("País de origen") and not str(row[0]).__contains__(" "):
            results.append((data['Year'].iloc[0], data['Month'].iloc[0], '', *row.values[:-2], 'V'))

    try:
        conn = sqlite3.connect('DDBB\\amazon.db')
        cursor = conn.cursor()
        # delete all data if exists
        delete_query = f"DELETE FROM sales WHERE Year = :year AND Month = :month and Type='V';"
        cursor.execute(delete_query, {"year": year, "month": month})
        # insert new data from results
        df_result = pd.DataFrame(results, columns=['Year', 'Month', 'Country', 'description', 'import', 'iva', 'Total', 'Type'])
        df_result.to_sql("sales", conn, if_exists="append", index=False)
        conn.close()
        messagebox.showinfo(title="Upload Data", message="Sales Data Loaded")
    except Exception as e:
        print(f"An error occurred: {e}")